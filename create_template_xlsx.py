import openpyxl
from openpyxl.styles import Font, Alignment

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active

# Add header row with bold formatting
ws['A1'] = 'STT'
ws['B1'] = 'MSSV'
ws['C1'] = 'Họ tên'
ws['D1'] = 'Giới tính'
ws['E1'] = 'Ngày sinh'

# Apply bold formatting to header
for cell in ['A1', 'B1', 'C1', 'D1', 'E1']:
    ws[cell].font = Font(bold=True)
    ws[cell].alignment = Alignment(horizontal='center')

# Add sample data with Vietnamese characters
ws['A2'] = 1
ws['B2'] = 'SV001'
ws['C2'] = 'Nguyễn Văn A'
ws['D2'] = 'Nam'
ws['E2'] = '15/01/2005'

ws['A3'] = 2
ws['B3'] = 'SV002'
ws['C3'] = 'Trần Thị B'
ws['D3'] = 'Nữ'
ws['E3'] = '20/03/2005'

ws['A4'] = 3
ws['B4'] = 'SV003'
ws['C4'] = 'Lê Văn C'
ws['D4'] = 'Nam'
ws['E4'] = '10/12/2004'

ws['A5'] = 4
ws['B5'] = 'SV004'
ws['C5'] = 'Phạm Thị D'
ws['D5'] = 'Nữ'
ws['E5'] = '25/05/2005'

ws['A6'] = 5
ws['B6'] = 'SV005'
ws['C6'] = 'Hoàng Văn E'
ws['D6'] = 'Nam'
ws['E6'] = '08/07/2005'

# Set column widths for better readability
ws.column_dimensions['A'].width = 10
ws.column_dimensions['B'].width = 15
ws.column_dimensions['C'].width = 25
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 15

# Save file
wb.save('static/templates/enroll_template.xlsx')
print('Created enroll_template.xlsx successfully with format: STT, MSSV, Họ tên, Giới tính, Ngày sinh')
